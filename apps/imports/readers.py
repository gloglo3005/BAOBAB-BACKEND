"""Lecture des fichiers d'import (CSV et XLSX).

Cette couche ne sait rien des livres : elle transforme un fichier en lignes
« colonne canonique -> texte ». Le sens des colonnes (catalogue, plus tard
membres ou espaces de coworking) reste dans le module qui l'appelle.
"""
import csv
import io
import os
import re
import unicodedata
from dataclasses import dataclass, field

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = (".csv", ".xlsx")

# Au-delà de ce nombre de lignes vides d'affilée, on arrête de lire (une feuille
# Excel mise en forme jusqu'à la ligne 1 000 000 ne doit pas bloquer le serveur).
MAX_BLANK_RUN = 1000


class ReadError(Exception):
    """Le fichier ne peut pas être lu ; le message est destiné à l'agent."""


@dataclass
class Row:
    line: int
    values: dict


@dataclass
class Table:
    rows: list = field(default_factory=list)
    columns: list = field(default_factory=list)
    ignored_columns: list = field(default_factory=list)


def normalize_header(text):
    """« Catégorie », « categorie » et « CATEGORIE » donnent la même clé."""
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = text.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stocke 1956 sous la forme 1956.0
    return " ".join(str(value).split())


def _build_table(raw_rows, aliases, required):
    """raw_rows : itérable de (numéro de ligne, liste de cellules)."""
    iterator = iter(raw_rows)

    header = None
    for line, cells in iterator:
        if any(_cell_text(cell) for cell in cells):
            header = cells
            break
    if header is None:
        raise ReadError("Le fichier est vide.")

    mapping, ignored = {}, []
    for index, name in enumerate(header):
        label = _cell_text(name)
        if not label:
            continue
        canonical = aliases.get(normalize_header(label))
        if canonical is None:
            ignored.append(label)
            continue
        if canonical in mapping.values():
            raise ReadError(f"La colonne « {canonical} » apparaît deux fois dans le fichier.")
        mapping[index] = canonical

    missing = [name for name in required if name not in mapping.values()]
    if missing:
        raise ReadError(
            "Colonne obligatoire absente : " + ", ".join(missing)
            + ". Vérifiez la première ligne du fichier ou téléchargez le modèle."
        )

    table = Table(columns=list(mapping.values()), ignored_columns=ignored)
    blank_run = 0
    for line, cells in iterator:
        values = {
            canonical: _cell_text(cells[index]) if index < len(cells) else ""
            for index, canonical in mapping.items()
        }
        if not any(values.values()):
            blank_run += 1
            if blank_run > MAX_BLANK_RUN:
                break
            continue
        blank_run = 0
        table.rows.append(Row(line=line, values=values))
        if len(table.rows) > MAX_ROWS:
            raise ReadError(
                f"Le fichier dépasse {MAX_ROWS:,} lignes".replace(",", " ")
                + " (limite d'un import). Découpez-le en plusieurs fichiers."
            )

    if not table.rows:
        raise ReadError("Le fichier ne contient aucune ligne de données.")
    return table


def _read_csv(data, aliases, required):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ReadError(
            "Le fichier CSV doit être enregistré en UTF-8. Dans Excel : "
            "Enregistrer sous, puis « CSV UTF-8 (délimité par des virgules) »."
        ) from None
    if "\x00" in text:
        raise ReadError("Ce fichier n'est pas un CSV lisible.")

    first_line = text.split("\n", 1)[0]
    delimiter = max((";", ",", "\t"), key=first_line.count)
    if not first_line.count(delimiter):
        delimiter = ","

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)

    def rows():
        try:
            for cells in reader:
                yield reader.line_num, cells
        except csv.Error as error:
            raise ReadError(f"Le fichier CSV est mal formé ({error}).") from error

    return _build_table(rows(), aliases, required)


def _read_xlsx(data, aliases, required):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ReadError(
            "La lecture des fichiers XLSX n'est pas disponible sur ce serveur "
            "(bibliothèque openpyxl manquante)."
        ) from None

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
    except Exception as error:
        raise ReadError("Impossible de lire ce fichier XLSX (fichier abîmé ?).") from error

    def rows():
        try:
            for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
                yield number, list(cells)
        except Exception as error:
            raise ReadError("Impossible de lire ce fichier XLSX (fichier abîmé ?).") from error

    try:
        return _build_table(rows(), aliases, required)
    finally:
        workbook.close()


def read_table(data, filename, aliases, required=()):
    """Lit un fichier CSV ou XLSX (première feuille) et renvoie un `Table`.

    `aliases` : {en-tête normalisé -> nom de colonne canonique}.
    `required` : colonnes canoniques obligatoires.
    """
    extension = os.path.splitext(filename or "")[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ReadError("Format non pris en charge : envoyez un fichier .csv ou .xlsx.")
    if len(data) > MAX_BYTES:
        raise ReadError("Le fichier dépasse 5 Mo.")
    if not data:
        raise ReadError("Le fichier est vide.")
    if extension == ".xlsx":
        return _read_xlsx(data, aliases, required)
    return _read_csv(data, aliases, required)
